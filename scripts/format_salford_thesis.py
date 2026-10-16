"""Apply the public Salford formatting guide to the generated Word thesis."""

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "thesis" / "final_thesis.docx"


def field_paragraph(code: str):
    paragraph = OxmlElement("w:p")
    for kind, value in (
        ("begin", None),
        ("instruction", code),
        ("separate", None),
        ("text", "Update field in Word to display the contents."),
        ("end", None),
    ):
        run = OxmlElement("w:r")
        if kind in {"begin", "separate", "end"}:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        elif kind == "instruction":
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = f" {value} "
        else:
            node = OxmlElement("w:t")
            node.text = value
        run.append(node)
        paragraph.append(run)
    return paragraph


def main() -> None:
    doc = Document(PATH)
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(3)
        section.right_margin = Cm(3)

    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        style.font.name = "Arial"
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5
        if style.name in {"Normal", "Body Text", "First Paragraph", "Compact", "Block Text"}:
            style.font.size = Pt(12)
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif style.name.startswith("Heading"):
            style.font.size = Pt(14 if style.name == "Heading 1" else 12)
            style.paragraph_format.keep_with_next = True

    for name in ("Figure Caption", "Table Caption", "Contents Title"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(12)
    doc.styles["Figure Caption"].paragraph_format.line_spacing = 1.5
    doc.styles["Table Caption"].paragraph_format.line_spacing = 1.5
    doc.styles["Table Caption"].paragraph_format.keep_with_next = True
    doc.styles["Contents Title"].font.bold = True
    doc.styles["Contents Title"].font.size = Pt(14)

    top_level = {
        "Declaration",
        "Contents",
        "List of figures",
        "List of tables",
        "Abbreviations",
        "Abstract",
        "References",
    }
    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()
        if value in top_level or value.startswith(("Chapter ", "Appendix ")):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.page_break_before = value not in {
                "Contents",
                "Chapter 1: Introduction",
            }
        if value == "Contents":
            paragraph.style = doc.styles["Contents Title"]
        elif re.match(r"Figure \d+\.\d+\.", value) and paragraph.style.name != "Compact":
            paragraph.style = doc.styles["Figure Caption"]
        elif re.match(r"Table [A-Z0-9]+\.\d+\.", value) and paragraph.style.name != "Compact":
            paragraph.style = doc.styles["Table Caption"]

    fields = {
        "Contents": 'TOC \\o "1-3" \\h \\z \\u',
        "List of figures": 'TOC \\h \\z \\t "Figure Caption,1"',
        "List of tables": 'TOC \\h \\z \\t "Table Caption,1"',
    }
    for heading, code in fields.items():
        paragraphs = list(doc.paragraphs)
        index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading)
        anchor = paragraphs[index]
        if heading != "Contents":
            for item in paragraphs[index + 1 :]:
                if item.style.name.startswith("Heading"):
                    break
                if item.text.strip():
                    item._element.getparent().remove(item._element)
        anchor._p.addnext(field_paragraph(code))

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if len(table.columns) == 5 and table.cell(0, 1).text == "Paired cases":
            table.cell(0, 1).text = "Cases"
            table.autofit = False
            widths = [0.32, 0.12, 0.17, 0.24, 0.15]
            for column, fraction in zip(table.columns, widths, strict=True):
                column.width = Cm(15 * fraction)
            for row in table.rows:
                for cell, fraction in zip(row.cells, widths, strict=True):
                    cell.width = Cm(15 * fraction)
        for row_number, row in enumerate(table.rows):
            row_props = row._tr.get_or_add_trPr()
            if row_number == 0:
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                row_props.append(repeat)
            cant_split = OxmlElement("w:cantSplit")
            row_props.append(cant_split)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.5
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.keep_with_next = row_number == 0
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(12)

    doc.save(PATH)
    print(f"Formatted {PATH}")


if __name__ == "__main__":
    main()
